"""C 域 Excel 导出工具（基于 openpyxl）。"""
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_workbook(title, headers, rows):
    """构造工作簿：表头加粗居中、冻结首行、自适应列宽。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row_index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col, value=value)

    ws.freeze_panes = "A2"
    for col in range(1, len(headers) + 1):
        width = len(str(headers[col - 1])) * 2 + 6
        for row in rows:
            if col - 1 < len(row) and row[col - 1] is not None:
                width = max(width, len(str(row[col - 1])) * 2 + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(width, 50)
    return wb


def xlsx_response(filename, title, headers, rows):
    """生成 Excel 下载响应；中文文件名按 RFC 5987 编码。"""
    wb = build_workbook(title, headers, rows)
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response
