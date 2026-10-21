"""判分与评分接口（C 域 #6-#11、#21）。

接口清单：
6.  GET    /api/exam/students/{id}/extracurricular/     查看学生课外学习材料
7.  GET    /api/exam/students/{id}/scores/              查看学生测评成绩
8.  POST   /api/exam/students/{id}/scores/daily/        为学生平时表现评分
9.  GET    /api/exam/scores/template/                   成绩导入模板下载
10. POST   /api/exam/scores/import/                     导入学生线下成绩
11. POST   /api/exam/scores/supplement/                 测评成绩补录
21. GET    /api/exam/students/{id}/scores/export/       导出学生所有成绩
"""
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from rest_framework.views import APIView

from apps.classroom.models import ClassRoom, ClassStudent
from apps.common.exceptions import ApiException
from apps.common.permissions import has_role
from apps.common.response import ApiResponse, ErrorCode
from apps.exam import services
from apps.exam.exports import xlsx_response
from apps.exam.models import AnswerSheet, ExtracurricularUpload, PerformanceScore
from apps.questionbank.models import Course
from apps.users.models import User

SUBJECT_NAME_MAP = {
    "music": "音乐", "音乐": "music",
    "art": "美术", "美术": "art",
}


def _student_scores(student):
    """学生的测评成绩（已提交答卷）与平时表现评分。"""
    sheets = (
        AnswerSheet.objects.filter(student=student, is_submitted=True, is_deleted=False)
        .select_related("paper")
        .order_by("-submitted_at")
    )
    performances = (
        PerformanceScore.objects.filter(student=student, is_deleted=False)
        .select_related("course", "klass", "scorer")
        .order_by("-created_at")
    )
    return sheets, performances


def _student_score_rows(student):
    """导出用行数据。"""
    sheets, performances = _student_scores(student)
    rows = []
    for sheet in sheets:
        rows.append([
            "测评", sheet.paper.title, sheet.paper.get_subject_display(),
            float(sheet.total_score), float(sheet.paper.total_score),
            sheet.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if sheet.submitted_at else "",
            "",
        ])
    for item in performances:
        rows.append([
            "平时表现", "平时表现评分", item.course.name if item.course else "",
            float(item.score), 100,
            item.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            item.comment,
        ])
    return rows


class StudentExtracurricularView(APIView):
    """GET /api/exam/students/{id}/extracurricular/ 查看学生课外学习材料。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def get(self, request, student_id):
        student = services.get_student_or_404(request.user, student_id)
        uploads = ExtracurricularUpload.objects.filter(student=student, is_deleted=False)
        data = [
            {
                "id": item.id,
                "title": item.title,
                "description": item.description,
                "image": request.build_absolute_uri(item.image.url) if item.image else "",
                "created_at": item.created_at,
            }
            for item in uploads
        ]
        return ApiResponse(data={"student_id": student.id, "list": data, "total": len(data)})


class StudentScoresView(APIView):
    """GET /api/exam/students/{id}/scores/ 查看学生测评成绩。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def get(self, request, student_id):
        student = services.get_student_or_404(request.user, student_id)
        sheets, performances = _student_scores(student)

        exam_list = [
            {
                "sheet_id": sheet.id,
                "paper_id": sheet.paper_id,
                "title": sheet.paper.title,
                "subject": sheet.paper.subject,
                "subject_display": sheet.paper.get_subject_display(),
                "ptype": sheet.paper.ptype,
                "score": float(sheet.total_score),
                "total_score": float(sheet.paper.total_score),
                "submitted_at": sheet.submitted_at,
            }
            for sheet in sheets
        ]
        performance_list = [
            {
                "id": item.id,
                "course": item.course.name if item.course else "",
                "klass": item.klass.name if item.klass else "",
                "score": float(item.score),
                "comment": item.comment,
                "scorer": item.scorer.real_name or item.scorer.username if item.scorer else "",
                "created_at": item.created_at,
            }
            for item in performances
        ]
        averages = {}
        for row in exam_list:
            averages.setdefault(row["subject_display"], []).append(row["score"])
        summary = {name: round(sum(vals) / len(vals), 2) for name, vals in averages.items() if vals}

        return ApiResponse(data={
            "student": {
                "id": student.id,
                "username": student.username,
                "real_name": student.real_name,
                "student_no": student.student_no,
            },
            "exams": exam_list,
            "performances": performance_list,
            "subject_average": summary,
        })


class StudentDailyScoreView(APIView):
    """POST /api/exam/students/{id}/scores/daily/ 为学生平时表现评分。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def post(self, request, student_id):
        student = services.get_student_or_404(request.user, student_id)

        raw_score = request.data.get("score")
        if raw_score in (None, ""):
            raise ApiException("请填写评分", ErrorCode.INVALID_PARAM, 400)
        try:
            score = Decimal(str(raw_score))
        except (InvalidOperation, ValueError):
            raise ApiException("评分格式不正确", ErrorCode.INVALID_PARAM, 400)
        if score < 0 or score > 100:
            raise ApiException("评分需在 0-100 之间", ErrorCode.INVALID_PARAM, 400)

        course = None
        course_id = request.data.get("course_id")
        if course_id:
            course = Course.objects.filter(id=course_id).first()
            if course is None:
                raise ApiException("学科不存在", ErrorCode.NOT_FOUND, 404)

        klass = None
        klass_id = request.data.get("klass_id")
        if klass_id:
            klass = ClassRoom.objects.filter(id=klass_id).first()
            if klass is None:
                raise ApiException("班级不存在", ErrorCode.NOT_FOUND, 404)
        else:
            klass_id = (
                ClassStudent.objects.filter(student=student, status=services.APPROVED)
                .values_list("klass_id", flat=True)
                .first()
            )
            klass = ClassRoom.objects.filter(id=klass_id).first() if klass_id else None

        record = PerformanceScore.objects.create(
            student=student,
            klass=klass,
            course=course,
            scorer=request.user,
            score=score,
            comment=request.data.get("comment", ""),
        )
        return ApiResponse(
            data={
                "id": record.id,
                "student_id": student.id,
                "score": float(record.score),
                "course": course.name if course else "",
                "comment": record.comment,
            },
            message="评分成功",
        )


class ScoreTemplateView(APIView):
    """GET /api/exam/scores/template/ 成绩导入模板下载。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def get(self, request):
        headers = ["学籍号", "姓名", "班级编号", "学科", "分数", "评语"]
        rows = [
            ["20260101", "张三", "C2601", "音乐", 88, "课堂表现积极"],
            ["20260102", "李四", "C2601", "美术", 92, "作品完成度高"],
        ]
        return xlsx_response("成绩导入模板.xlsx", "成绩导入", headers, rows)


class ScoreImportView(APIView):
    """POST /api/exam/scores/import/ 导入学生线下成绩（xlsx）。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def post(self, request):
        upload = request.FILES.get("file")
        if upload is None:
            raise ApiException("请上传 Excel 文件（字段名 file）", ErrorCode.INVALID_PARAM, 400)
        try:
            workbook = load_workbook(upload, data_only=True)
        except Exception:
            raise ApiException("文件解析失败，请使用模板文件", ErrorCode.INVALID_PARAM, 400)

        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        allowed = services.scoped_students(request.user)

        success, errors = 0, []
        for index, row in enumerate(rows, start=2):
            if not row or all(value in (None, "") for value in row):
                continue
            values = list(row) + [None] * (6 - len(row))
            student_no, real_name, class_code, subject, score_value, comment = values[:6]

            student = None
            if student_no:
                student = allowed.filter(student_no=str(student_no).strip()).first()
            if student is None and real_name:
                student = allowed.filter(real_name=str(real_name).strip()).first()
            if student is None:
                errors.append({"row": index, "error": "学生不存在或不在管理范围内"})
                continue

            course = None
            if subject:
                code = SUBJECT_NAME_MAP.get(str(subject).strip())
                if code is None:
                    errors.append({"row": index, "error": "学科仅支持 音乐/美术"})
                    continue
                course = Course.objects.filter(code=code).first() or Course.objects.filter(
                    name="音乐" if code == "music" else "美术"
                ).first()

            try:
                score = Decimal(str(score_value))
            except (InvalidOperation, ValueError):
                errors.append({"row": index, "error": "分数格式不正确"})
                continue
            if score < 0 or score > 100:
                errors.append({"row": index, "error": "分数需在 0-100 之间"})
                continue

            klass = None
            if class_code:
                klass = ClassRoom.objects.filter(code=str(class_code).strip()).first()

            PerformanceScore.objects.create(
                student=student,
                klass=klass,
                course=course,
                scorer=request.user,
                score=score,
                comment=str(comment or ""),
            )
            success += 1

        return ApiResponse(
            data={"total": success + len(errors), "success": success, "failed": len(errors), "errors": errors},
            message="导入完成",
        )


class ScoreSupplementView(APIView):
    """POST /api/exam/scores/supplement/ 测评成绩补录（学校/区县/市州）。"""

    permission_classes = [has_role(User.ROLE_SCHOOL, User.ROLE_DISTRICT, User.ROLE_CITY)]

    def post(self, request):
        items = request.data.get("items")
        payloads = items if isinstance(items, list) else [request.data]
        if not payloads:
            raise ApiException("请提供补录数据", ErrorCode.INVALID_PARAM, 400)

        results, errors = [], []
        for index, item in enumerate(payloads):
            student_id = item.get("student_id")
            student_no = item.get("student_no")
            student = None
            if student_id:
                student = User.objects.filter(id=student_id, role=User.ROLE_STUDENT).first()
            elif student_no:
                student = User.objects.filter(student_no=student_no, role=User.ROLE_STUDENT).first()
            if student is None or not services.scoped_students(request.user).filter(id=student.id).exists():
                errors.append({"index": index, "error": "学生不存在或不在管理范围内"})
                continue

            try:
                score = Decimal(str(item.get("score")))
            except (InvalidOperation, ValueError, TypeError):
                errors.append({"index": index, "error": "分数格式不正确"})
                continue
            if score < 0 or score > 100:
                errors.append({"index": index, "error": "分数需在 0-100 之间"})
                continue

            course = Course.objects.filter(id=item.get("course_id")).first() if item.get("course_id") else None
            klass = ClassRoom.objects.filter(id=item.get("klass_id")).first() if item.get("klass_id") else None
            comment = str(item.get("comment") or "")
            record = PerformanceScore.objects.create(
                student=student,
                klass=klass,
                course=course,
                scorer=request.user,
                score=score,
                comment=f"测评补录：{comment}" if comment else "测评补录",
            )
            results.append({
                "id": record.id,
                "student_id": student.id,
                "student_no": student.student_no,
                "score": float(record.score),
            })

        return ApiResponse(
            data={"success": len(results), "failed": len(errors), "results": results, "errors": errors},
            message="补录完成",
        )


class StudentScoresExportView(APIView):
    """GET /api/exam/students/{id}/scores/export/ 导出学生所有成绩。"""

    permission_classes = [has_role(User.ROLE_TEACHER, User.ROLE_SCHOOL)]

    def get(self, request, student_id):
        student = services.get_student_or_404(request.user, student_id)
        headers = ["类型", "名称", "学科", "得分", "满分", "时间", "评语"]
        rows = _student_score_rows(student)
        filename = f"{student.real_name or student.username}_成绩.xlsx"
        return xlsx_response(filename, "学生成绩", headers, rows)
